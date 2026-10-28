"""Excel contents, report parity, filtering and access checks on real report data."""

from __future__ import annotations

import io
import re
import uuid
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from urllib.parse import unquote

import pytest
import pytest_asyncio
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import delete
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import Settings
from app.db.models import (
    AssessmentJob,
    AssessmentResult,
    Assignment,
    AssignmentAccess,
    AssignmentRevision,
    SourcePrompt,
    Submission,
    SubmissionAttempt,
    TranscriptionRevision,
    User,
    UserRole,
)
from app.services.report_export import (
    EXCEL_DATETIME_FORMAT,
    SCORE_SOURCE_LABELS,
    STATUS_LABELS,
    XLSX_MEDIA_TYPE,
    build_workbook,
    content_disposition,
    local_datetime,
    safe_cell,
)
from tests.catalog.conftest import (
    VALID_RUBRIC,
    create_term,
    create_user,
    login_as,
    seed_classroom_with_student,
)

EXPORTED_AT = datetime(2026, 12, 31, 18, 30, tzinfo=timezone.utc)
SUBMITTED_AT = datetime(2026, 10, 27, 23, 5, tzinfo=timezone.utc)


def score_payload(awards: list[float]) -> dict:
    return {
        "totalScore": sum(awards),
        "maxScore": 4,
        "evaluations": [
            {"criterionId": criterion["id"], "criterionName": criterion["name"],
             "awardedScore": award, "maxScore": criterion["weight"]}
            for criterion, award in reversed(list(zip(VALID_RUBRIC["criteria"], awards, strict=True)))
        ],
    }


@pytest_asyncio.fixture
async def report_data(
    db_session: AsyncSession, teacher_user: User, student_user: User,
) -> dict:
    term = await create_term(db_session, ended=False)
    classroom = await seed_classroom_with_student(db_session, teacher_user, student_user, term=term)
    prompt = SourcePrompt(owner_teacher_id=teacher_user.id, title="Đề xuất Excel", status="READY")
    db_session.add(prompt)
    await db_session.flush()
    assignment = Assignment(
        teacher_id=teacher_user.id, classroom_id=classroom.id, source_prompt_id=prompt.id,
        title="Bài kiểm tra tiếng Việt", code="EXCEL-01", status="PUBLISHED",
        assignment_rubric=VALID_RUBRIC, max_score=4, published_at=SUBMITTED_AT,
    )
    db_session.add(assignment)
    await db_session.flush()
    revision = AssignmentRevision(
        assignment_id=assignment.id, code=assignment.code, title=assignment.title,
        genre="NLXH", rubric=VALID_RUBRIC, snapshot={"rubric": VALID_RUBRIC}, status="PUBLISHED",
    )
    db_session.add(revision)
    await db_session.flush()
    assignment.published_revision_id = revision.id
    student_user.display_name = "=2+2"
    students = [student_user] + [
        User(login_id=f"excel_{index:03d}", email=f"excel_{index}@example.test",
             display_name=f"Học sinh {index:03d}", password_hash="x", role=UserRole.STUDENT)
        for index in range(1, 53)
    ]
    db_session.add_all(students[1:])
    await db_session.flush()
    db_session.add_all([
        AssignmentAccess(assignment_revision_id=revision.id, student_id=student.id)
        for student in students
    ])
    results = []
    for index, student in enumerate(students[:4]):
        submission = Submission(assignment_revision_id=revision.id, student_id=student.id)
        db_session.add(submission)
        await db_session.flush()
        submitted_at = SUBMITTED_AT - timedelta(days=1 if index == 2 else 0)
        confirmed_at = submitted_at + timedelta(minutes=10)
        attempt = SubmissionAttempt(
            submission_id=submission.id, assignment_revision_id=revision.id, attempt_number=1,
            state="OCR_REVIEW" if index == 3 else "READY_TO_GRADE", page_count=1,
            submitted_at=submitted_at, official_submitted_at=None if index == 3 else confirmed_at,
        )
        db_session.add(attempt)
        await db_session.flush()
        submission.current_attempt_id = attempt.id
        if index == 3:
            continue
        tx = TranscriptionRevision(attempt_id=attempt.id, revision=1, status="CONFIRMED")
        db_session.add(tx)
        await db_session.flush()
        attempt.confirmed_revision_id = tx.id
        job = AssessmentJob(
            submission_id=submission.id, attempt_id=attempt.id, transcription_revision_id=tx.id,
            assignment_revision_id=revision.id, student_id=student.id, status="READY",
        )
        db_session.add(job)
        await db_session.flush()
        ai_payload = score_payload([[1, 0, 0, 0], [0, 0, 0, 0], [2, 0.25, 0.25, 0]][index])
        teacher_payload = score_payload([2, 0, 0, 0]) if index == 0 else None
        published = index != 1
        published_data = score_payload([2.5, 0.5, 0, 0] if index == 0 else [1, 0, 0, 0]) if published else None
        total = (teacher_payload or ai_payload)["totalScore"]
        result = AssessmentResult(
            job_id=job.id, submission_id=submission.id, attempt_id=attempt.id,
            transcription_revision_id=tx.id, assignment_revision_id=revision.id, student_id=student.id,
            ai_payload=ai_payload, teacher_payload=teacher_payload, published_payload=published_data,
            ai_total_score=ai_payload["totalScore"], total_score=total, max_score=4,
            score_step=0.25, score_scale10=total * 2.5,
            is_published=published, published_at=confirmed_at if published else None,
        )
        db_session.add(result)
        await db_session.flush()
        submission.current_assessment_job_id = job.id
        submission.current_assessment_result_id = result.id
        submission.current_published_result_id = result.id if published else None
        results.append(result)
    await db_session.commit()
    return {"assignment": assignment, "revision": revision, "students": students,
            "results": results, "classroom": classroom, "term": term}


async def read_report_and_export(client: AsyncClient, assignment_id: uuid.UUID, params: dict):
    base = f"/api/teacher/reports/{assignment_id}"
    report_response = await client.get(base, params={**params, "pageSize": 50})
    assert report_response.status_code == 200, report_response.text
    report = report_response.json()
    students = list(report["students"]["items"])
    for page in range(2, report["students"]["totalPages"] + 1):
        response = await client.get(base, params={**params, "pageSize": 50, "page": page})
        assert response.status_code == 200
        students.extend(response.json()["students"]["items"])
    response = await client.get(base + "/export.xlsx", params=params)
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content))
    return report, students, response, workbook


@pytest.mark.parametrize("mode", [None, "official", "provisional"])
async def test_export_matches_report_scores_all_pages_and_statistics(
    client: AsyncClient, report_data: dict, monkeypatch, mode: str | None,
) -> None:
    monkeypatch.setattr("app.services.report_export.utcnow", lambda: EXPORTED_AT)
    report, students, response, workbook = await read_report_and_export(
        client, report_data["assignment"].id, {"mode": mode} if mode else {},
    )
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    disposition = response.headers["content-disposition"]
    assert 'attachment; filename="bang-diem-EXCEL-01-20270101.xlsx"' in disposition
    assert "filename*=UTF-8''" in disposition
    assert "bảng-điểm-EXCEL-01-20270101.xlsx" in unquote(disposition)
    assert workbook.sheetnames == ["Bảng điểm", "Thống kê"]
    grades = workbook["Bảng điểm"]
    rows = list(grades.values)
    assert len(rows) == 54  # 53 students, including the second page and unscored students.
    assert len(students) == report["students"]["total"] == 53
    assert rows[0][:9] == (
        "STT", "Họ tên", "Mã đăng nhập", "Trạng thái", "Nộp ảnh lúc", "Xác nhận lúc",
        "Điểm theo thang barem", "Điểm thang 10", "Nguồn điểm",
    )
    assert rows[0][9:13] == (
        "Nội dung (tối đa 2.5)", "Lập luận (tối đa 0.75)",
        "Diễn đạt (tối đa 0.5)", "Bố cục (tối đa 0.25)",
    )
    assert rows[0][13:] == ("Nộp muộn", "Hạn hiệu lực", "Bản", "Điểm bản 1", "Chênh so với bản 1")
    by_student = {str(result.student_id): result for result in report_data["results"]}
    for index, (row, student) in enumerate(zip(rows[1:], students, strict=True), start=1):
        assert row[:4] == (index, safe_cell(student["name"]), student["code"], STATUS_LABELS[student["status"]])
        assert row[4:9] == (
            local_datetime(student["ocrSubmittedAt"]), local_datetime(student["officialSubmittedAt"]),
            student["score"], student["scoreScale10"], SCORE_SOURCE_LABELS.get(student["scoreSource"]),
        )
        if student["score"] is None:
            assert row[9:13] == (None,) * 4
        else:
            result = by_student[student["studentId"]]
            payload = result.published_payload if report["mode"] == "official" else result.teacher_payload or result.ai_payload
            awards = {item["criterionId"]: item["awardedScore"] for item in payload["evaluations"]}
            assert row[9:13] == tuple(awards[item["id"]] for item in VALID_RUBRIC["criteria"])
        first = student["firstPublishedScore"]
        delta = student["score"] - first if student["score"] is not None and first is not None else None
        assert row[13:] == (
            "Có" if student["isLate"] else "Không", local_datetime(student["effectiveDueAt"]),
            student["revisionRound"], first, delta,
        )
    assert rows[1][4] == datetime(2026, 10, 28, 6, 5)
    assert rows[1][5] == datetime(2026, 10, 28, 6, 15)
    assert grades["E2"].number_format == grades["F2"].number_format == EXCEL_DATETIME_FORMAT
    assert rows[1][6:8] == ((2, 5) if mode == "provisional" else (3, 7.5))
    assert grades["B2"].value == "'=2+2" and grades["B2"].data_type == "s"
    assert grades.freeze_panes == "D2"
    assert grades.auto_filter.ref == "A1:R54"
    summary_rows = list(workbook["Thống kê"].values)
    summary = {row[0]: row[1] for row in summary_rows}
    assert summary["Tên bài"] == report["assignment"]["title"]
    assert summary["Lớp"] == report["assignment"]["classroomName"]
    assert summary["Thời điểm xuất"] == datetime(2027, 1, 1, 1, 30)
    assert summary["Chế độ điểm"] == ("Tạm tính" if report["mode"] == "provisional" else "Chính thức")
    for label, key in {
        "Số học sinh được giao": "assigned", "Số học sinh đã nộp chính thức": "officialSubmitted",
        "Số học sinh đã công bố": "published", "Số học sinh có điểm": "scoredCount",
        "Điểm trung bình theo thang barem": "meanScore", "Điểm trung bình thang 10": "meanScoreScale10",
        "Điểm thấp nhất thang 10": "minScoreScale10", "Điểm cao nhất thang 10": "maxScoreScale10",
    }.items():
        assert summary[label] == pytest.approx(report["stats"][key])
    distribution_rows = {row[0]: row[1:] for row in summary_rows[-5:]}
    for bucket in report["stats"]["distribution"]:
        assert distribution_rows[bucket["range"]] == pytest.approx((bucket["count"], bucket["rate"]))
    workbook.close()


@pytest.mark.parametrize("filters", [
    {"search": "Học sinh", "status": "not_submitted", "sort": "score_desc"},
    {"search": "EXCEL_002", "status": "published", "mode": "official"},
    {"status": "pending", "mode": "provisional", "sort": "score_asc"},
    {"mode": "provisional", "from": "2026-10-27", "to": "2026-10-28"},
    {"mode": "official", "from": "2026-10-27T20:00:00Z", "to": "2026-10-28T01:00:00+00:00"},
    {"search": "Không có học sinh này"},
])
async def test_export_uses_same_filters_and_sort_as_report(client: AsyncClient, report_data: dict, filters: dict) -> None:
    report, students, _response, workbook = await read_report_and_export(client, report_data["assignment"].id, filters)
    rows = list(workbook["Bảng điểm"].values)[1:]
    assert len(rows) == len(students) == report["students"]["total"]
    for row, student in zip(rows, students, strict=True):
        assert row[1:3] == (safe_cell(student["name"]), student["code"])
        assert row[6:8] == (student["score"], student["scoreScale10"])
        if student["score"] is None:
            assert row[9:13] == (None,) * 4
    workbook.close()


async def test_export_date_aliases_and_student_scope(client: AsyncClient, report_data: dict) -> None:
    assignment_id = report_data["assignment"].id
    student_id = str(report_data["students"][2].id)
    params = {"studentId": student_id, "mode": "official", "from": "2026-10-27", "to": "2026-10-28"}
    _report, students, _response, workbook = await read_report_and_export(client, assignment_id, params)
    response = await client.get(f"/api/teacher/reports/{assignment_id}/export.xlsx", params={
        "studentId": student_id, "mode": "official", "dateFrom": params["from"], "dateTo": params["to"],
    })
    assert response.status_code == 200
    alias_workbook = load_workbook(io.BytesIO(response.content))
    assert list(alias_workbook["Bảng điểm"].values) == list(workbook["Bảng điểm"].values)
    assert len(students) == 1 and students[0]["score"] is None
    assert list(alias_workbook["Bảng điểm"].values)[1][9:13] == (None,) * 4
    workbook.close()
    alias_workbook.close()


async def test_export_defaults_to_provisional_without_published_scores(
    client: AsyncClient, db_session: AsyncSession, report_data: dict,
) -> None:
    for result in report_data["results"]:
        result.is_published = False
    await db_session.commit()
    report, _students, _response, workbook = await read_report_and_export(client, report_data["assignment"].id, {})
    assert report["mode"] == report["suggestedMode"] == "provisional"
    assert workbook["Bảng điểm"]["G2"].value == 2
    assert workbook["Thống kê"]["B4"].value == "Tạm tính"
    workbook.close()


async def test_export_empty_assignment_keeps_rubric_headers(
    client: AsyncClient, db_session: AsyncSession, report_data: dict,
) -> None:
    await db_session.execute(delete(AssignmentAccess).where(AssignmentAccess.assignment_revision_id == report_data["revision"].id))
    await db_session.commit()
    report, _students, _response, workbook = await read_report_and_export(client, report_data["assignment"].id, {})
    assert report["students"]["total"] == 0
    assert workbook["Bảng điểm"].max_row == 1
    assert workbook["Bảng điểm"].max_column == 18
    assert workbook["Thống kê"]["B6"].value == 0
    assert workbook["Thống kê"]["B11"].value is None
    workbook.close()


async def test_export_is_readable_for_ended_term(
    client: AsyncClient, db_session: AsyncSession, report_data: dict,
) -> None:
    report_data["term"].ends_on = datetime(2020, 1, 1).date()
    report_data["term"].starts_on = datetime(2019, 1, 1).date()
    await db_session.commit()
    response = await client.get(f"/api/teacher/reports/{report_data['assignment'].id}/export.xlsx")
    assert response.status_code == 200


@pytest.mark.parametrize("params, code", [
    ({"mode": "invalid"}, "STATS_MODE"),
    ({"dateFrom": "not-a-date"}, "VALIDATION"),
    ({"to": "2026-10-99"}, "VALIDATION"),
])
async def test_export_rejects_invalid_filters(client: AsyncClient, report_data: dict, params: dict, code: str) -> None:
    response = await client.get(f"/api/teacher/reports/{report_data['assignment'].id}/export.xlsx", params=params)
    assert response.status_code == 400
    assert response.json()["error"]["code"] == code


@pytest.mark.parametrize("viewer", ["other_teacher", "student", "anonymous"])
async def test_export_access_matches_existing_report(
    client: AsyncClient, db_session: AsyncSession, settings: Settings, report_data: dict, viewer: str,
) -> None:
    if viewer == "other_teacher":
        user = await create_user(db_session, settings, client, role=UserRole.TEACHER, prefix="other")
        await login_as(client, settings, db_session, user)
    elif viewer == "student":
        await login_as(client, settings, db_session, report_data["students"][0])
    else:
        client.cookies.clear()
    base = f"/api/teacher/reports/{report_data['assignment'].id}"
    expected = await client.get(base)
    response = await client.get(base + "/export.xlsx")
    assert response.status_code == expected.status_code == {"other_teacher": 404, "student": 403, "anonymous": 401}[viewer]
    assert response.json() == expected.json()


@pytest.mark.parametrize("value", ["=SUM(1,2)", "+1+2", "-1+2", "@SUM(A1)", "\t=1+2", "\r=1+2"])
async def test_workbook_sanitizes_every_text_cell(client: AsyncClient, report_data: dict, value: str) -> None:
    # Exercise persisted workbook cells, including dynamic headers and summary text.
    response = await client.get(f"/api/teacher/reports/{report_data['assignment'].id}")
    assert response.status_code == 200
    report = response.json()
    report["assignment"].update(title=value, classroomName=value)
    student = deepcopy(report["students"]["items"][0])
    student.update(name=value, code=value, status=value, scoreSource=value)
    content = build_workbook(report, [student], [{"id": "C1", "name": value, "weight": 4}], {}, exported_at=EXPORTED_AT)
    workbook = load_workbook(io.BytesIO(content))
    for cell in ("B2", "C2", "D2", "I2"):
        assert workbook["Bảng điểm"][cell].value == "'" + value
        assert workbook["Bảng điểm"][cell].data_type == "s"
    assert workbook["Bảng điểm"]["J1"].value == "'" + value + " (tối đa 4)"
    for cell in ("B2", "B3"):
        assert workbook["Thống kê"][cell].value == "'" + value
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    assert safe_cell(-2.5) == -2.5 and safe_cell(0) == 0
    workbook.close()


def test_download_filename_handles_unicode_and_header_characters() -> None:
    header = content_disposition('ĐỀ-01/"\r\n', EXPORTED_AT)
    header.encode("ascii")
    assert "\r" not in header and "\n" not in header
    assert re.search(r'filename="bang-diem-[A-Za-z0-9_-]+-20270101.xlsx"', header)
    assert "bảng-điểm-ĐỀ-01-20270101.xlsx" in unquote(header)
