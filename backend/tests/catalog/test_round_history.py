"""Round visibility, published snapshots, report baselines and export parity."""

from __future__ import annotations

import io
import uuid
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.http import utcnow
from app.db.models import AssessmentResult, AssignmentAccess, Submission, UserRole
from app.db.session import get_session_factory
from app.services.report_export import EXCEL_DATETIME_FORMAT, local_datetime
from app.services.round_compare import compare_rounds
from tests.catalog.conftest import HEADERS, create_user, login_as
from tests.catalog.test_deadlines import (
    deadline_clock as deadline_clock,
    make_assignment as make_assignment,
)
from tests.catalog.test_revision_rounds import (
    confirm_round,
    first_publication,
    grade_round,
    publish_round_without_grading,
    result_target,
    student_path,
    teacher_path,
)


def student_rounds_path(assignment):
    return student_path(assignment) + "/assessment/rounds"


def teacher_rounds_path(assignment, student):
    return f"/api/teacher/assignments/{assignment['id']}/students/{student.id}/rounds"


async def perfect_publication(client, assignment, student, teacher, db_session, settings):
    await grade_round(settings)
    await login_as(client, settings, db_session, teacher)
    path = teacher_path(assignment, student)
    detail = (await client.get(path)).json()
    reviewed = await client.patch(path, headers=HEADERS, json={
        **result_target(detail["assessment"]["currentResult"]),
        "evaluations": [
            {"criterionId": item["criterionId"], "awardedScore": item["maxScore"]}
            for item in detail["assessment"]["working"]["evaluations"]
        ],
        "overallComment": "Nhận xét bản viết lại đã công bố.",
        "improvementGuidance": ["Tiếp tục phát huy."],
    })
    assert reviewed.status_code == 200, reviewed.text
    return await publish_round_without_grading(client, assignment, student, teacher, db_session, settings)


async def report(client, assignment, **params):
    response = await client.get(f"/api/teacher/reports/{assignment['id']}", params={"mode": "official", **params})
    assert response.status_code == 200, response.text
    return response.json()


@pytest.mark.asyncio
async def test_history_keeps_snapshots_and_teacher_sees_pending_round(
    client, make_assignment, teacher_user, student_user, db_session, settings,
):
    assignment, _first, pub1 = await first_publication(
        client, make_assignment, teacher_user, student_user, db_session, settings, maxRevisions=2,
    )
    response = await client.get(student_rounds_path(assignment))
    assert response.status_code == 200, response.text
    first = response.json()["rounds"][0]
    assert first["round"] == 1 and first["resultId"] == pub1["id"]
    assert first["visibleToStudent"] and first["isPublished"]
    assert response.json()["comparison"] is None
    await login_as(client, settings, db_session, teacher_user)
    initial = await report(client, assignment)
    assert initial["students"]["items"][0]["firstPublishedScore"] == first["totalScore"]
    assert initial["students"]["items"][0]["publishedRound"] == 1
    assert (await report(client, assignment, status="revised"))["students"]["total"] == 0

    await login_as(client, settings, db_session, student_user)
    started = await client.post(student_path(assignment) + "/submission/start-revision", headers=HEADERS)
    assert started.status_code == 200
    await confirm_round(client, assignment, settings)
    student_pending = (await client.get(student_rounds_path(assignment))).json()
    assert student_pending["rounds"] == [first]
    assert student_pending["revisionRound"] == 2 and student_pending["currentPublishedRound"] == 1
    await login_as(client, settings, db_session, teacher_user)
    teacher_pending = await client.get(teacher_rounds_path(assignment, student_user))
    assert teacher_pending.status_code == 200, teacher_pending.text
    pending = teacher_pending.json()["rounds"][-1]
    assert pending["round"] == 2 and pending["jobStatus"] == "QUEUED"
    assert pending["resultId"] is None and pending["totalScore"] is None
    assert not pending["visibleToStudent"] and pending["publishedAt"] is None
    revised = await report(client, assignment, status="revised")
    row = revised["students"]["items"][0]
    assert row["revisionRound"] == 2 and row["publishedRound"] == 1
    assert row["firstPublishedScore"] == first["totalScore"]
    # A new round adds history, never another contribution to score aggregates.
    for key in ("scoredCount", "meanScore", "meanScoreScale10", "minScoreScale10", "maxScoreScale10", "distribution"):
        assert revised["stats"][key] == initial["stats"][key]

    await grade_round(settings)
    ready = (await client.get(teacher_rounds_path(assignment, student_user))).json()["rounds"][-1]
    assert ready["round"] == 2 and ready["jobStatus"] == "READY"
    assert ready["resultId"] and ready["totalScore"] is not None and not ready["visibleToStudent"]
    pub2 = await publish_round_without_grading(client, assignment, student_user, teacher_user, db_session, settings)
    published = (await client.get(student_rounds_path(assignment))).json()
    assert [item["round"] for item in published["rounds"]] == [1, 2]
    old, current = published["rounds"]
    assert old == {**first, "isPublished": False, "status": "superseded"}
    assert current["resultId"] == pub2["id"] and current["isPublished"]
    assert published["currentPublishedRound"] == 2
    assert published["comparison"] == compare_rounds(old, current)
    assert published["comparison"]["totalScoreDelta"] == 0
    # Later unannounced working edits must not leak into either round snapshot.
    async with get_session_factory()() as session:
        result = await session.get(AssessmentResult, uuid.UUID(pub1["id"]))
        result.teacher_payload = {"totalScore": 999, "overallComment": "Ghi chú chưa công bố"}
        result.total_score = 999
        await session.commit()
    assert (await client.get(student_rounds_path(assignment))).json() == published
    await login_as(client, settings, db_session, teacher_user)
    final = await report(client, assignment)
    final_row = final["students"]["items"][0]
    assert final_row["firstPublishedScore"] == first["totalScore"]
    assert final_row["firstPublishedScoreScale10"] == first["scoreScale10"]
    assert final_row["publishedRound"] == final_row["revisionRound"] == 2
    for key in ("scoredCount", "meanScore", "meanScoreScale10", "minScoreScale10", "maxScoreScale10", "distribution"):
        assert final["stats"][key] == initial["stats"][key]

    await login_as(client, settings, db_session, student_user)
    assert (await client.post(student_path(assignment) + "/submission/start-revision", headers=HEADERS)).status_code == 200
    await confirm_round(client, assignment, settings)
    await perfect_publication(client, assignment, student_user, teacher_user, db_session, settings)
    third = (await client.get(student_rounds_path(assignment))).json()
    assert [item["round"] for item in third["rounds"]] == [1, 2, 3]
    assert third["currentPublishedRound"] == 3
    assert third["comparison"] == compare_rounds(third["rounds"][1], third["rounds"][2])
    # Legacy replacements without a superseded marker are not student-visible.
    # In particular, their scores must not leak through a comparison across a gap.
    async with get_session_factory()() as session:
        second = await session.get(AssessmentResult, uuid.UUID(pub2["id"]))
        second.superseded_at = None
        await session.commit()
    gap = (await client.get(student_rounds_path(assignment))).json()
    assert [item["round"] for item in gap["rounds"]] == [1, 3]
    assert gap["comparison"] is None
    await login_as(client, settings, db_session, teacher_user)
    teacher_history = (await client.get(teacher_rounds_path(assignment, student_user))).json()
    assert [item["round"] for item in teacher_history["rounds"]] == [1, 2, 3]
    assert not teacher_history["rounds"][1]["visibleToStudent"]
    assert teacher_history["comparison"]["previousRound"] == 2


@pytest.mark.asyncio
async def test_manual_unpublish_hides_current_round_and_comparison_but_preserves_teacher_history(
    client, make_assignment, teacher_user, student_user, db_session, settings,
):
    assignment, _first, _pub1 = await first_publication(
        client, make_assignment, teacher_user, student_user, db_session, settings, maxRevisions=1,
    )
    assert (await client.post(student_path(assignment) + "/submission/start-revision", headers=HEADERS)).status_code == 200
    await confirm_round(client, assignment, settings)
    pub2 = await perfect_publication(client, assignment, student_user, teacher_user, db_session, settings)
    before = (await client.get(student_rounds_path(assignment))).json()
    assert before["rounds"][-1]["totalScore"] == before["rounds"][-1]["maxScore"]
    assert before["comparison"]["totalScoreDelta"] > 0
    await login_as(client, settings, db_session, teacher_user)
    response = await client.post(teacher_path(assignment, student_user) + "/unpublish", headers=HEADERS, json=result_target(pub2))
    assert response.status_code == 200, response.text
    hidden = (await client.get(teacher_rounds_path(assignment, student_user))).json()
    assert [item["visibleToStudent"] for item in hidden["rounds"]] == [True, False]
    assert hidden["rounds"][-1]["status"] == "unpublished"
    assert hidden["rounds"][-1]["overallComment"] == "Nhận xét bản viết lại đã công bố."
    assert len(hidden["rounds"]) == 2  # No duplicate pending entry for a hidden result.
    await login_as(client, settings, db_session, student_user)
    response = (await client.get(student_rounds_path(assignment))).json()
    assert response["rounds"] == before["rounds"][:1]
    assert response["currentPublishedRound"] is None and response["comparison"] is None


@pytest.mark.asyncio
async def test_same_round_corrections_and_legacy_history_stay_hidden_from_student(
    client, make_assignment, teacher_user, student_user, db_session, settings,
):
    assignment, first, pub1 = await first_publication(
        client, make_assignment, teacher_user, student_user, db_session, settings, maxRevisions=1,
    )
    await login_as(client, settings, db_session, teacher_user)
    assert (await client.post(teacher_path(assignment, student_user) + "/grant-reopen", headers=HEADERS)).status_code == 200
    await login_as(client, settings, db_session, student_user)
    assert (await client.post(student_path(assignment) + "/ocr/reopen", headers=HEADERS)).status_code == 200
    assert (await client.post(student_path(assignment) + "/ocr/confirm", headers=HEADERS)).status_code == 200
    corrected = await perfect_publication(client, assignment, student_user, teacher_user, db_session, settings)
    async with get_session_factory()() as session:
        sub = await session.get(Submission, uuid.UUID(first["submissionId"]))
        assert "reason" not in sub.published_history[-1]
    visible = (await client.get(student_rounds_path(assignment))).json()
    assert [item["resultId"] for item in visible["rounds"]] == [corrected["id"]]
    assert visible["comparison"] is None
    await login_as(client, settings, db_session, teacher_user)
    history = (await client.get(teacher_rounds_path(assignment, student_user))).json()["rounds"]
    assert {item["resultId"]: item["visibleToStudent"] for item in history} == {pub1["id"]: False, corrected["id"]: True}
    row = (await report(client, assignment))["students"]["items"][0]
    assert row["revisionRound"] == row["publishedRound"] == 1
    assert row["firstPublishedScore"] == row["score"]
    assert (await report(client, assignment, status="revised"))["students"]["total"] == 0
    # Hiding the corrected round 1 still leaves its published score as the baseline.
    hidden = await client.post(
        teacher_path(assignment, student_user) + "/unpublish", headers=HEADERS, json=result_target(corrected),
    )
    assert hidden.status_code == 200, hidden.text
    hidden_row = (await report(client, assignment))["students"]["items"][0]
    assert hidden_row["publishedRound"] is None and hidden_row["score"] is None
    assert hidden_row["firstPublishedScore"] == row["score"]
    assert hidden_row["firstPublishedScoreScale10"] == row["scoreScale10"]


@pytest.mark.asyncio
async def test_round_history_requires_recipient_ownership_and_roles(
    client, make_assignment, teacher_user, student_user, db_session, settings,
):
    assignment = await make_assignment(maxRevisions=1)
    teacher_url = teacher_rounds_path(assignment, student_user)
    empty = await client.get(teacher_url)
    assert empty.status_code == 200 and empty.json()["rounds"] == []
    await login_as(client, settings, db_session, student_user)
    empty = await client.get(student_rounds_path(assignment))
    assert empty.status_code == 200 and empty.json()["rounds"] == []
    assert empty.json()["comparison"] is None
    assert (await client.get(teacher_url)).status_code == 403
    other_student = await create_user(db_session, settings, client, role=UserRole.STUDENT, prefix="outsider")
    await login_as(client, settings, db_session, other_student)
    denied = await client.get(student_rounds_path(assignment))
    assert denied.status_code == 403 and denied.json()["error"]["code"] == "ACCESS_DENIED"
    await login_as(client, settings, db_session, teacher_user)
    assert (await client.get(student_rounds_path(assignment))).status_code == 403
    assert (await client.get(teacher_rounds_path(assignment, other_student))).status_code == 403
    other_teacher = await create_user(db_session, settings, client, role=UserRole.TEACHER, prefix="other_teacher")
    await login_as(client, settings, db_session, other_teacher)
    assert (await client.get(teacher_url)).status_code == 404
    client.cookies.clear()
    assert (await client.get(student_rounds_path(assignment))).status_code == 401
    assert (await client.get(teacher_url)).status_code == 401


@pytest.mark.asyncio
async def test_revoked_recipient_cannot_read_published_rounds(
    client, make_assignment, teacher_user, student_user, db_session, settings,
):
    assignment, _first, _pub1 = await first_publication(
        client, make_assignment, teacher_user, student_user, db_session, settings, maxRevisions=1,
    )
    async with get_session_factory()() as session:
        grant = (await session.execute(select(AssignmentAccess).where(AssignmentAccess.student_id == student_user.id))).scalar_one()
        grant.revoked_at = utcnow()
        await session.commit()
    denied = await client.get(student_rounds_path(assignment))
    assert denied.status_code == 403 and denied.json()["error"]["code"] == "ACCESS_DENIED"


@pytest.mark.asyncio
async def test_revision_excel_columns_use_report_fields_and_latest_selected_score(
    client, make_assignment, teacher_user, student_user, db_session, settings, deadline_clock,
):
    assignment, _first, _pub1 = await first_publication(
        client, make_assignment, teacher_user, student_user, db_session, settings, maxRevisions=1,
    )
    await login_as(client, settings, db_session, teacher_user)
    before = await report(client, assignment)
    row1 = before["students"]["items"][0]
    first_score = row1["score"]
    # Deadline and extension columns use A's fields and the Excel local timezone.
    due = (deadline_clock[0] - timedelta(minutes=1)).replace(microsecond=0)
    response = await client.put(
        f"/api/teacher/assignments/{assignment['id']}/students/{student_user.id}/due-extension",
        headers=HEADERS, json={"dueAt": due.isoformat()},
    )
    assert response.status_code == 200, response.text
    export_url = f"/api/teacher/reports/{assignment['id']}/export.xlsx"
    response = await client.get(export_url, params={"mode": "official", "status": "late"})
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["Bảng điểm"]
    values = dict(zip(next(sheet.values), list(sheet.values)[1], strict=True))
    assert values["Nộp muộn"] == "Có"
    assert values["Hạn hiệu lực"] == local_datetime(due)
    due_cell = next(cell for cell in sheet[2] if cell.column == list(values).index("Hạn hiệu lực") + 1)
    assert due_cell.number_format == EXCEL_DATETIME_FORMAT
    assert values["Bản"] == 1 and values["Điểm bản 1"] == first_score
    assert values["Chênh so với bản 1"] == 0
    workbook.close()
    await login_as(client, settings, db_session, student_user)
    assert (await client.post(student_path(assignment) + "/submission/start-revision", headers=HEADERS)).status_code == 200
    await confirm_round(client, assignment, settings)
    await perfect_publication(client, assignment, student_user, teacher_user, db_session, settings)
    await login_as(client, settings, db_session, teacher_user)
    current = await report(client, assignment, status="revised")
    row = current["students"]["items"][0]
    response = await client.get(export_url, params={"mode": "official", "status": "revised"})
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["Bảng điểm"]
    assert sheet.max_row == 2
    values = dict(zip(next(sheet.values), list(sheet.values)[1], strict=True))
    assert values["Bản"] == 2 and values["Nộp muộn"] == "Không"
    assert values["Điểm bản 1"] == row["firstPublishedScore"] == first_score
    assert values["Chênh so với bản 1"] == row["score"] - first_score > 0
    assert values["Điểm theo thang barem"] == row["score"]
    assert current["stats"]["scoredCount"] == before["stats"]["scoredCount"] == 1
    assert current["stats"]["meanScore"] == row["score"]
    assert current["stats"]["meanScoreScale10"] == row["scoreScale10"]
    assert sum(bucket["count"] for bucket in current["stats"]["distribution"]) == 1
    workbook.close()
