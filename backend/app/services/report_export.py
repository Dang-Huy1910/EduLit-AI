"""Excel exports of the existing class report, including every matching row."""

from __future__ import annotations

import io
import re
import uuid
from datetime import datetime, timezone
from typing import Any
from urllib.parse import quote
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.http import utcnow
from app.db.models import User
from app.services.assessment_scoring import AssessmentInvalidOutput, rubric_from_revision
from app.services.report_service import ReportService
from app.services.round_compare import score_delta
from app.services.stats_engine import MODE_PROVISIONAL, parse_time_bound, select_score

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LOCAL_TIMEZONE = ZoneInfo("Asia/Ho_Chi_Minh")
EXCEL_DATETIME_FORMAT = "dd/mm/yyyy hh:mm"

# Keep in sync with frontend/src/features/assessment/labels.ts (rowStatusLabel).
STATUS_LABELS = {
    "not_submitted": "Chưa nộp",
    "processing": "Đang xử lý",
    "ocr_review": "Đang rà OCR",
    "ocr_failed": "Lỗi OCR",
    "grading": "Đang chấm",
    "awaiting_review": "Chờ rà soát",
    "published": "Đã công bố",
    "grading_failed": "Lỗi chấm",
    "revising": "Đang sửa bản chép",
    "new_version_grading": "Bản mới đang chấm",
    "new_version_review": "Bản mới chờ công bố",
    "new_version_failed": "Bản mới lỗi chấm",
}
SCORE_SOURCE_LABELS = {"published": "Đã công bố", "teacher": "Giáo viên", "ai": "AI"}


def safe_cell(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def local_datetime(value: datetime | str | None) -> datetime | None:
    if not value:
        return None
    moment = datetime.fromisoformat(value.replace("Z", "+00:00")) if isinstance(value, str) else value
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=timezone.utc)
    # Excel datetimes cannot carry timezone information.
    return moment.astimezone(LOCAL_TIMEZONE).replace(tzinfo=None)


def grade_headers(criteria: list[dict[str, Any]]) -> list[str]:
    return [
        "STT", "Họ tên", "Mã đăng nhập", "Trạng thái", "Nộp ảnh lúc", "Xác nhận lúc",
        "Điểm theo thang barem", "Điểm thang 10", "Nguồn điểm",
        *[f"{item['name']} (tối đa {item['weight']:g})" for item in criteria],
        "Nộp muộn", "Hạn hiệu lực", "Bản", "Điểm bản 1", "Chênh so với bản 1",
    ]


def grade_row(
    index: int,
    student: dict[str, Any],
    payload: dict[str, Any] | None,
    criteria: list[dict[str, Any]],
) -> list[Any]:
    """Export the selected report score and the published round 1 baseline."""
    awards = {
        str(item.get("criterionId") or ""): item.get("awardedScore")
        for item in (payload or {}).get("evaluations") or []
        if isinstance(item, dict)
    }
    return [
        index,
        student["name"],
        student["code"],
        STATUS_LABELS.get(student["status"], student["status"]),
        local_datetime(student["ocrSubmittedAt"]),
        local_datetime(student["officialSubmittedAt"]),
        student["score"],
        student["scoreScale10"],
        SCORE_SOURCE_LABELS.get(student["scoreSource"], student["scoreSource"]),
        *[awards.get(str(item["id"])) for item in criteria],
        "Có" if student.get("isLate") else "Không",
        local_datetime(student.get("effectiveDueAt")),
        student.get("revisionRound", 1),
        student.get("firstPublishedScore"),
        score_delta(student.get("firstPublishedScore"), student["score"]),
    ]


def _append_row(sheet: Worksheet, values: list[Any]) -> None:
    sheet.append([safe_cell(value) for value in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, datetime):
            cell.number_format = EXCEL_DATETIME_FORMAT


def build_workbook(
    report: dict[str, Any],
    students: list[dict[str, Any]],
    criteria: list[dict[str, Any]],
    payloads: dict[str, dict[str, Any]],
    *,
    exported_at: datetime,
) -> bytes:
    workbook = Workbook()
    grades = workbook.active
    grades.title = "Bảng điểm"
    _append_row(grades, grade_headers(criteria))
    for index, student in enumerate(students, start=1):
        # A date filter can leave a student visible with no selected score.
        payload = payloads.get(student["studentId"]) if student["score"] is not None else None
        _append_row(grades, grade_row(index, student, payload, criteria))
    grades.freeze_panes = "D2"
    grades.auto_filter.ref = grades.dimensions
    grades.row_dimensions[1].height = 36
    widths = [6, 28, 22, 28, 22, 22, 24, 20, 20] + [28] * len(criteria) + [16, 22, 10, 18, 24]
    for index, width in enumerate(widths, start=1):
        grades.column_dimensions[get_column_letter(index)].width = width

    summary = workbook.create_sheet("Thống kê")
    stats = report["stats"]
    assignment = report["assignment"]
    mode_label = "Tạm tính" if report["mode"] == MODE_PROVISIONAL else "Chính thức"
    for row in [
        ["Nội dung", "Giá trị"],
        ["Tên bài", assignment["title"]],
        ["Lớp", assignment["classroomName"]],
        ["Chế độ điểm", mode_label],
        ["Thời điểm xuất", local_datetime(exported_at)],
        ["Số học sinh được giao", stats["assigned"]],
        ["Số học sinh đã nộp chính thức", stats["officialSubmitted"]],
        ["Số học sinh đã công bố", stats["published"]],
        ["Số học sinh có điểm", stats["scoredCount"]],
        ["Điểm trung bình theo thang barem", stats["meanScore"]],
        ["Điểm trung bình thang 10", stats["meanScoreScale10"]],
        ["Điểm thấp nhất thang 10", stats["minScoreScale10"]],
        ["Điểm cao nhất thang 10", stats["maxScoreScale10"]],
        ["Khoảng điểm thang 10", "Số học sinh", "Tỷ lệ"],
    ]:
        _append_row(summary, row)
    for bucket in stats["distribution"]:
        _append_row(summary, [bucket["range"], bucket["count"], bucket["rate"]])
        summary.cell(summary.max_row, 3).number_format = "0.0%"
    summary.column_dimensions["A"].width = 40
    summary.column_dimensions["B"].width = 52
    summary.column_dimensions["C"].width = 16
    summary.freeze_panes = "B2"

    for sheet in (grades, summary):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E2E8F0")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def content_disposition(code: str, exported_at: datetime) -> str:
    stamp = exported_at.astimezone(LOCAL_TIMEZONE).strftime("%Y%m%d")
    ascii_code = re.sub(r"[^A-Za-z0-9_-]+", "-", code).strip("-") or "bai-tap"
    unicode_code = re.sub(r"[^\w-]+", "-", code).strip("-") or "bai-tap"
    ascii_name = f"bang-diem-{ascii_code}-{stamp}.xlsx"
    unicode_name = f"bảng-điểm-{unicode_code}-{stamp}.xlsx"
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(unicode_name, safe='')}"


async def export_class_report(
    service: ReportService,
    teacher: User,
    assignment_id: uuid.UUID,
    *,
    mode: str | None = None,
    search: str = "",
    status: str = "all",
    sort: str = "default",
    date_from: str | None = None,
    date_to: str | None = None,
    student_id: uuid.UUID | None = None,
) -> tuple[bytes, str]:
    # class_report applies ownership checks and resolves the suggested score mode.
    report = await service.class_report(
        teacher, assignment_id, mode=mode, search=search, status=status, sort=sort,
        date_from=date_from, date_to=date_to, student_id=student_id, page_size=50,
    )
    assignment, revision, classroom = await service._owned(teacher, assignment_id)
    papers = await service.stats.papers_for_assignment(assignment, revision, classroom, student_id=student_id)
    students = list(report["students"]["items"])
    start = parse_time_bound(date_from, end=False)
    end = parse_time_bound(date_to, end=True)
    # The shared row helper caps page_size at 50. Reuse it for every remaining page.
    for page in range(2, report["students"]["totalPages"] + 1):
        result = service.stats.student_rows_for_assignment(
            papers, mode=report["mode"], search=search, status=status, sort=sort,
            start=start, end=end, page=page, page_size=50,
        )
        students.extend(result["items"])
    if papers:
        rubric = papers[0].rubric
    else:
        try:
            rubric = rubric_from_revision(revision.rubric, revision.snapshot)
        except AssessmentInvalidOutput:
            rubric = {}
    payloads = {}
    for paper in papers:
        selected = select_score(paper, report["mode"])
        if selected:
            payloads[str(paper.student.id)] = selected.payload
    exported_at = utcnow()
    content = build_workbook(report, students, rubric.get("criteria") or [], payloads, exported_at=exported_at)
    return content, content_disposition(assignment.code, exported_at)
